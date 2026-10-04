#!/usr/bin/env python3
"""LeadScope website-sales enrichment pipeline.

Reads an XLSX/CSV lead list, discovers public business emails from websites,
analyzes the public website, scores sales opportunity, and generates a
personalized outreach draft. It NEVER sends email.
"""
import argparse, asyncio, csv, json, os, re, sqlite3, sys, time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, asdict
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook, Workbook

EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
SOCIAL_HOSTS = ("instagram.com", "facebook.com", "tiktok.com", "x.com", "linkedin.com")
CONTACT_HINTS = ("contact", "kontakt", "impressum", "imprint", "about", "studio")
BOOKING_HINTS = ("book", "booking", "termin", "appointment", "reserv", "calendly", "fresha", "treatwell")
PORTFOLIO_HINTS = ("gallery", "portfolio", "artists", "arbeiten", "tattoo")
CTA_HINTS = ("book", "booking", "termin", "appointment", "kontakt", "contact", "anfrage")

@dataclass
class Analysis:
    final_url: str = ""
    http_status: int = 0
    https: bool = False
    title: str = ""
    meta_description: str = ""
    emails: str = ""
    email_source: str = ""
    contact_url: str = ""
    booking_detected: bool = False
    booking_evidence: str = ""
    mobile_viewport: bool = False
    cta_detected: bool = False
    social_links: str = ""
    page_words: int = 0
    image_count: int = 0
    broken_fetch: bool = False
    analysis_notes: str = ""
    website_score: int = 0
    sales_opportunity: int = 0
    opportunity_tier: str = "SKIP"
    pitch_angle: str = ""
    outreach_subject: str = ""
    outreach_draft: str = ""


def norm_url(url):
    if not url: return ""
    url = str(url).strip()
    if not url: return ""
    if not url.startswith(("http://", "https://")): url = "https://" + url
    return url


def fetch(url, timeout=15):
    headers = {"User-Agent": "LeadScope/1.0 (+website audit; contact owner before outreach)"}
    return requests.get(url, headers=headers, timeout=timeout, allow_redirects=True)


def same_domain(a,b):
    return urlparse(a).netloc.lower().removeprefix("www.") == urlparse(b).netloc.lower().removeprefix("www.")


def extract_links(base, soup):
    out=[]
    for a in soup.find_all("a", href=True):
        href=urljoin(base,a.get("href"))
        if same_domain(base, href): out.append((a.get_text(" ",strip=True), href))
    return out


def analyze_site(url):
    a=Analysis()
    url=norm_url(url)
    if not url:
        a.analysis_notes="No website supplied"
        return a
    try:
        r=fetch(url)
        a.final_url=r.url; a.http_status=r.status_code; a.https=r.url.lower().startswith("https://")
        soup=BeautifulSoup(r.text, "html.parser")
        a.title=(soup.title.get_text(" ",strip=True) if soup.title else "")[:200]
        md=soup.find("meta", attrs={"name":"description"})
        a.meta_description=(md.get("content","") if md else "")[:500]
        text=soup.get_text(" ",strip=True)
        a.page_words=len(text.split()); a.image_count=len(soup.find_all("img"))
        viewport=soup.find("meta", attrs={"name":"viewport"})
        a.mobile_viewport=bool(viewport and "width" in (viewport.get("content","").lower()))
        links=extract_links(r.url,soup)
        all_text=" ".join([t for t,_ in links]).lower()+" "+text.lower()
        booking=[]; contact=[]
        for t,h in links:
            low=(t+" "+h).lower()
            if any(x in low for x in BOOKING_HINTS): booking.append(h)
            if any(x in low for x in CONTACT_HINTS): contact.append(h)
        a.booking_detected=bool(booking); a.booking_evidence=booking[0] if booking else ""
        a.contact_url=contact[0] if contact else ""
        a.cta_detected=any(x in all_text for x in CTA_HINTS)
        socials=[]
        for _,h in [(t,h) for t,h in links]:
            if any(host in urlparse(h).netloc.lower() for host in SOCIAL_HOSTS): socials.append(h)
        a.social_links=", ".join(dict.fromkeys(socials))[:1000]
        emails=EMAIL_RE.findall(r.text)
        # Check likely contact pages too, but cap requests.
        checked={r.url}; candidates=[]
        for _,h in links:
            if h not in checked and any(x in (h.lower()) for x in CONTACT_HINTS): candidates.append(h)
        for h in candidates[:2]:
            try:
                rr=fetch(h,timeout=10); checked.add(rr.url)
                emails += EMAIL_RE.findall(rr.text)
            except Exception: pass
        emails=[e.lower() for e in emails if not e.lower().endswith((".png",".jpg",".jpeg",".webp"))]
        # Filter obvious non-contact technical addresses.
        emails=[e for e in dict.fromkeys(emails) if not any(x in e for x in ("example.com","sentry","wixpress","wordpress"))]
        a.emails=", ".join(emails[:5]); a.email_source="website/contact page" if emails else ""
        notes=[]
        if not a.mobile_viewport: notes.append("no mobile viewport detected")
        if not a.meta_description: notes.append("no meta description detected")
        if not a.cta_detected: notes.append("weak/no obvious CTA")
        if not a.booking_detected: notes.append("no obvious booking/appointment path")
        if not a.contact_url: notes.append("no obvious contact page")
        if a.page_words < 150: notes.append("thin homepage copy")
        a.analysis_notes="; ".join(notes)
    except Exception as e:
        a.broken_fetch=True; a.analysis_notes=f"Fetch failed: {type(e).__name__}: {e}"
    return a


def score(a: Analysis, rating=None, reviews=None, owner=None):
    # Website opportunity: deficiencies increase opportunity; active/business signals add confidence.
    weak=0
    if not a.https: weak+=10
    if not a.mobile_viewport: weak+=12
    if not a.meta_description: weak+=4
    if not a.cta_detected: weak+=10
    if not a.booking_detected: weak+=14
    if not a.contact_url: weak+=5
    if a.page_words < 150: weak+=8
    if a.image_count < 5: weak+=8
    if a.broken_fetch: return 0, "SKIP"
    if rating is not None:
        try:
            if float(rating)>=4.5: weak+=10
            elif float(rating)>=4.2: weak+=7
            elif float(rating)>=4.0: weak+=4
        except: pass
    if reviews is not None:
        try:
            n=int(float(reviews)); weak += min(8, n//50)
        except: pass
    if a.emails: weak+=8
    if owner: weak+=3
    s=max(0,min(100,weak))
    tier="A" if s>=65 else "B" if s>=45 else "C" if s>=25 else "SKIP"
    return s,tier


def choose_pitch(a):
    if a.broken_fetch: return "skip"
    if not a.mobile_viewport: return "mobile experience"
    if not a.booking_detected: return "booking/conversion flow"
    if not a.cta_detected: return "clearer booking CTA"
    if a.image_count < 5: return "portfolio presentation"
    if a.page_words < 150: return "clearer studio positioning"
    if not a.meta_description: return "search presentation"
    return "overall conversion-focused redesign"


def draft(biz, owner, a):
    name=owner.split("/")[0].strip() if owner else ""
    greeting=f"Hallo {name}," if name else "Hallo {biz}-Team,"
    angle=choose_pitch(a)
    subject=f"Kurze Idee für die Website von {biz}"
    observation={
        "mobile experience":"mir ist aufgefallen, dass die mobile Darstellung noch nicht so stark wirkt, wie es bei einem Studio mit euren Arbeiten möglich wäre",
        "booking/conversion flow":"mir ist aufgefallen, dass der Weg von euren Arbeiten zu einer konkreten Terminanfrage nicht ganz so direkt ist, wie er sein könnte",
        "clearer booking CTA":"mir ist aufgefallen, dass eine klare Handlungsaufforderung zur Termin-Anfrage nicht sofort im Vordergrund steht",
        "portfolio presentation":"mir ist aufgefallen, dass eure Arbeiten online noch stärker als visuelles Portfolio präsentiert werden könnten",
        "clearer studio positioning":"mir ist aufgefallen, dass ein neuer Besucher nicht sofort alles Wichtige über Studio, Artists und nächsten Schritt erfährt",
        "search presentation":"mir ist aufgefallen, dass eure Website bei der Darstellung in der Suche noch Potenzial hat",
        "overall conversion-focused redesign":"mir ist aufgefallen, dass man euren starken Studio-Auftritt online noch deutlich moderner und conversion-orientierter darstellen könnte",
    }[angle]
    body=(f"{greeting}\n\n"
          f"ich bin [DEIN NAME] von Giantic Studios. Ich bin gerade auf {biz} gestoßen und habe mir eure Website kurz angesehen. {observation}.\n\n"
          f"Ich baue Websites speziell für lokale Unternehmen und habe deshalb spontan eine Idee, wie man euren Auftritt neu aufbauen könnte – ohne eure Marke oder eure Arbeiten zu verlieren.\n\n"
          f"Wenn ihr möchtet, kann ich euch kostenlos einen kleinen Entwurf zeigen. Kein Vertrag und kein Verkaufsmeeting – einfach damit ihr seht, was ich meine.\n\n"
          f"Soll ich euch den Entwurf schicken?\n\n"
          f"Viele Grüße\n[DEIN NAME]\nGiantic Studios")
    return subject,body


def read_rows(path):
    p=Path(path)
    if p.suffix.lower()=='.csv':
        with p.open('r',encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))
    wb=load_workbook(p,read_only=True,data_only=True); ws=wb.active
    rows=list(ws.iter_rows(values_only=True)); headers=[str(x or '').strip() for x in rows[0]]
    return [dict(zip(headers,r)) for r in rows[1:]]


def write_xlsx(rows,path):
    wb=Workbook(); ws=wb.active; ws.title='LeadScope Pipeline'
    keys=list(rows[0].keys()) if rows else []
    ws.append(keys)
    for row in rows: ws.append([row.get(k,'') for k in keys])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        maxlen=min(60,max(len(str(c.value or '')) for c in col)+2); ws.column_dimensions[col[0].column_letter].width=maxlen
    wb.save(path)


def main():
    ap=argparse.ArgumentParser(); ap.add_argument('input'); ap.add_argument('-o','--output',default='leadscope_enriched.xlsx'); ap.add_argument('--workers',type=int,default=8); ap.add_argument('--limit',type=int); args=ap.parse_args()
    rows=read_rows(args.input); rows=rows[:args.limit] if args.limit else rows
    def one(row):
        a=analyze_site(row.get('website',''))
        s,t=score(a,row.get('google_rating'),row.get('google_review_count'),row.get('owner_name'))
        a.sales_opportunity=s; a.opportunity_tier=t; a.pitch_angle=choose_pitch(a)
        a.outreach_subject,a.outreach_draft=draft(row.get('business_name',''),row.get('owner_name',''),a)
        return row,a
    with ThreadPoolExecutor(max_workers=args.workers) as ex: results=list(ex.map(one,rows))
    out=[]
    for row,a in results:
        row=dict(row)
        d=asdict(a)
        row.update({f"ls_{k}":v for k,v in d.items()})
        out.append(row)
    write_xlsx(out,args.output)
    print(json.dumps({"input":args.input,"output":args.output,"processed":len(out),"tiers":{t:sum(1 for r in out if r.get('ls_opportunity_tier')==t) for t in ['A','B','C','SKIP']},"emails_found":sum(bool(r.get('ls_emails')) for r in out)},indent=2))

if __name__=='__main__': main()
